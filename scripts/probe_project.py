r"""
probe_project.py  --  inventory a scattered project and surface the files that likely
hold the three music-theory tables we need (17-note cents, interval names, Dang rules).

It does NOT read heavy data. For every Excel/CSV/code file it records only:
  - path, size, sheet names, row/column counts, and the header row + first 2 rows;
  - a "score" for how likely the file is a DEFINITION table (small, ~17 rows, or words
    like cent/vaziri/koron/interval/dang) versus a bulk OUTPUT table (153/246 rows).
Everything is written to  probe_report.txt  (small, text only) which you upload.

HOW TO RUN (Anaconda Prompt):

    conda activate music
    python probe_project.py  "D:\path\to\project_root_1"  "D:\another\folder"

You can pass one or several root folders (each in quotes). If you pass none, it scans
the current folder. Reading is read-only; nothing is modified.
"""

import os
import sys
import csv
import glob

MAXCELL = 60          # truncate long cell text
KEYWORDS = ["cent", "vaziri", "koron", "sori", "mojanab", "dang", "tetrachord",
            "interval", "pitch", "tuning", "temperament", "comma", "radian", "angle",
            "phase", "frequency", "shoor", "mahoor", "scale", "degree"]
CODE_EXT = (".py", ".r", ".m", ".ipynb", ".txt", ".json")
EXCEL_EXT = (".xlsx", ".xlsm", ".xls")
CSV_EXT = (".csv", ".tsv")


def kw_hits(text):
    t = str(text).lower()
    return [k for k in KEYWORDS if k in t]


def score_file(path, nrows, headertext):
    """Higher = more likely to be a small DEFINITION table we want."""
    s = 0
    base = os.path.basename(path).lower()
    s += 3 * len(set(kw_hits(base)))                 # keywords in filename
    s += len(set(kw_hits(headertext)))               # keywords in header
    if nrows is not None:
        if 15 <= nrows <= 30:
            s += 5                                    # ~17 pitch classes / ~24 intervals
        elif nrows in (0, None):
            pass
        elif nrows >= 120:
            s -= 3                                    # looks like bulk stimuli/output
    return s


def peek_excel(path, out):
    try:
        import openpyxl
    except ImportError:
        out.append("    (openpyxl not installed -> `pip install openpyxl`)")
        return -1, ""
    best_score = -99
    header_all = ""
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        out.append(f"    [could not open: {e}]")
        return -99, ""
    for sh in wb.sheetnames:
        ws = wb[sh]
        nrows, ncols = ws.max_row, ws.max_column
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            rows.append(row)
            if i >= 2:
                break
        header = rows[0] if rows else ()
        htext = " ".join("" if c is None else str(c) for c in header)
        header_all += " " + htext
        sc = score_file(path, nrows, htext + " " + sh)
        best_score = max(best_score, sc)
        out.append(f"    sheet '{sh}': {nrows} rows x {ncols} cols"
                   + (f"   <== keywords: {kw_hits(htext + ' ' + sh)}" if kw_hits(htext + ' ' + sh) else ""))
        for r in rows:
            cells = ["" if c is None else str(c)[:MAXCELL] for c in r]
            out.append("      | " + " | ".join(cells))
    wb.close()
    return best_score, header_all


def peek_csv(path, out):
    try:
        with open(path, "r", encoding="utf-8", errors="replace", newline="") as f:
            sample = f.read(4096)
        delim = "\t" if path.lower().endswith(".tsv") else ","
        rows = list(csv.reader(sample.splitlines(), delimiter=delim))[:3]
        # count total lines cheaply
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            nrows = sum(1 for _ in f)
    except Exception as e:
        out.append(f"    [could not open: {e}]"); return -99, ""
    header = rows[0] if rows else []
    htext = " ".join(header)
    out.append(f"    {nrows} rows"
               + (f"   <== keywords: {kw_hits(htext)}" if kw_hits(htext) else ""))
    for r in rows:
        out.append("      | " + " | ".join(str(c)[:MAXCELL] for c in r))
    return score_file(path, nrows, htext), htext


def peek_code(path, out):
    """Scan code/text files for lines that assign cents/angles/interval names/dang rules."""
    hits = []
    try:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            for i, line in enumerate(f, 1):
                low = line.lower()
                if any(k in low for k in ["cent", "vaziri", "koron", "sori", "mojanab",
                                          "dang", "tetrachord", "interval_name",
                                          "pi/6", "pi/4", "np.pi", "phase", "temperament"]):
                    hits.append((i, line.rstrip()[:120]))
                if len(hits) >= 12:
                    break
    except Exception:
        return -99, ""
    if hits:
        out.append(f"    {len(hits)} relevant line(s):")
        for ln, txt in hits:
            out.append(f"      L{ln}: {txt}")
    return (5 if hits else -1), " ".join(t for _, t in hits)


def main():
    roots = sys.argv[1:] or ["."]
    out = []
    candidates = []
    out.append("PROBE REPORT")
    out.append("roots: " + " | ".join(os.path.abspath(r) for r in roots))
    out.append("=" * 70)

    seen = 0
    for root in roots:
        for dirpath, _dirs, files in os.walk(root):
            for fn in sorted(files):
                ext = os.path.splitext(fn)[1].lower()
                if ext not in EXCEL_EXT + CSV_EXT + CODE_EXT:
                    continue
                path = os.path.join(dirpath, fn)
                try:
                    size = os.path.getsize(path)
                except OSError:
                    continue
                seen += 1
                block = [f"\n[{ext}] {path}  ({size/1024:.0f} KB)"]
                if ext in EXCEL_EXT:
                    sc, _ = peek_excel(path, block)
                elif ext in CSV_EXT:
                    sc, _ = peek_csv(path, block)
                else:
                    sc, _ = peek_code(path, block)
                out.extend(block)
                if sc >= 4:
                    candidates.append((sc, path))

    out.append("\n" + "=" * 70)
    out.append(f"scanned {seen} files")
    out.append("\nTOP CANDIDATE FILES (most likely to hold the definition tables):")
    for sc, path in sorted(candidates, reverse=True)[:25]:
        out.append(f"  score {sc:>3}  {path}")
    if not candidates:
        out.append("  (none scored high; upload probe_report.txt anyway and we'll look)")

    with open("probe_report.txt", "w", encoding="utf-8") as f:
        f.write("\n".join(out))
    print(f"Scanned {seen} files. Wrote probe_report.txt "
          f"({os.path.getsize('probe_report.txt')/1024:.0f} KB).")
    print("Upload probe_report.txt here.")
    if candidates:
        print("\nTop candidates:")
        for sc, path in sorted(candidates, reverse=True)[:10]:
            print(f"  score {sc:>3}  {path}")


if __name__ == "__main__":
    main()
